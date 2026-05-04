import asyncio
import csv
import io
import logging
import smtplib
import re
import json
import dns.resolver
from fastapi import Body, FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import HTMLResponse, StreamingResponse
from pathlib import Path
from pydantic import BaseModel
from openpyxl import load_workbook
import xlrd

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(title="Pyxis")


class NameInput(BaseModel):
    first: str
    last: str = ""
    middle: str = ""


class MultiNameRequest(BaseModel):
    domain: str
    names: list[NameInput]


class ValidateEmailsRequest(BaseModel):
    emails: list[str]


EMAIL_RE = re.compile(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}')

PATTERNS = [
    "{first}.{last}@{domain}",
    "{first}{last}@{domain}",
    "{f}.{last}@{domain}",
    "{first}.{l}@{domain}",
    "{f}{last}@{domain}",
    "{first}{l}@{domain}",
    "{first}@{domain}",
    "{last}@{domain}",
    "{f}.{l}@{domain}",
    "{f}{l}@{domain}",
    "{last}.{first}@{domain}",
    "{first}_{last}@{domain}",
    "{last}{first}@{domain}",
    "{first}.{middle}.{last}@{domain}",
    "{first}{middle}{last}@{domain}",
    "{f}{middle}{last}@{domain}",
]


DOMAIN_RE = re.compile(r"^(?!-)(?:[a-z0-9-]{1,63}\.)+[a-z]{2,63}$", re.IGNORECASE)


def normalize_domain(domain: str) -> str:
    value = domain.strip().lower()
    if not value:
        raise HTTPException(status_code=400, detail="Domain is required")
    if any(char.isspace() for char in value):
        raise HTTPException(status_code=400, detail="Domain cannot contain spaces")
    if "@" in value:
        raise HTTPException(status_code=400, detail="Enter only the domain, not an email address")
    if "://" in value:
        raise HTTPException(status_code=400, detail="Enter only the domain, not a URL")
    if "/" in value:
        raise HTTPException(status_code=400, detail="Enter only the domain, not a path")
    if "." not in value:
        raise HTTPException(status_code=400, detail="Domain must include a TLD, like .com or .ai")
    if not DOMAIN_RE.match(value):
        raise HTTPException(status_code=400, detail="Please enter a valid domain name with a TLD")
    return value


def extract_emails_from_text(text: str) -> list[str]:
    return EMAIL_RE.findall(text or "")


def dedupe_emails(emails: list[str]) -> list[str]:
    seen = set()
    result = []
    for email in emails:
        normalized = email.strip()
        if not normalized:
            continue
        key = normalized.lower()
        if key in seen:
            continue
        seen.add(key)
        result.append(normalized)
    return result


def extract_emails_from_csv_bytes(content: bytes) -> list[str]:
    text = content.decode("utf-8-sig", errors="ignore")
    reader = csv.reader(io.StringIO(text))
    found = []
    for row in reader:
        for cell in row:
            found.extend(extract_emails_from_text(str(cell)))
    return found


def extract_emails_from_xlsx_bytes(content: bytes) -> list[str]:
    found = []
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue
                found.extend(extract_emails_from_text(str(cell)))
    return found


def extract_emails_from_xls_bytes(content: bytes) -> list[str]:
    found = []
    workbook = xlrd.open_workbook(file_contents=content)
    for sheet in workbook.sheets():
        for row_idx in range(sheet.nrows):
            for cell in sheet.row_values(row_idx):
                if cell is None:
                    continue
                found.extend(extract_emails_from_text(str(cell)))
    return found


def extract_emails_from_file(filename: str, content: bytes) -> list[str]:
    lower_name = filename.lower()
    if lower_name.endswith(".csv"):
        return extract_emails_from_csv_bytes(content)
    if lower_name.endswith(".xlsx"):
        return extract_emails_from_xlsx_bytes(content)
    if lower_name.endswith(".xls"):
        return extract_emails_from_xls_bytes(content)
    if lower_name.endswith(".txt"):
        return extract_emails_from_text(content.decode("utf-8-sig", errors="ignore"))
    return extract_emails_from_text(content.decode("utf-8-sig", errors="ignore"))


def generate_patterns(first: str, last: str, domain: str, middle: str = "") -> list[str]:
    domain = domain.lower().strip()
    first = first.lower().strip()
    last = last.lower().strip()
    middle = middle.lower().strip()

    if not last:
        return [f"{first}@{domain}"]

    emails = set()
    for tmpl in PATTERNS:
        email = tmpl.format(
            first=first,
            last=last,
            f=first[0] if first else "",
            l=last[0] if last else "",
            middle=middle if middle else "",
            domain=domain,
        )
        # Clean up empty placeholder artifacts
        email = re.sub(r'\.@', '@', email)
        email = re.sub(r'\.\.', '.', email)
        email = email.lstrip('.')
        if re.match(r'^[^@]+@[^@]+\.[^@]+$', email):
            emails.add(email)
    return sorted(emails)


def normalize_person_name(person: NameInput) -> dict[str, str]:
    first = person.first.strip()
    if not first:
        return {}
    return {
        "first": first,
        "last": person.last.strip(),
        "middle": person.middle.strip(),
    }


def person_label(first: str, last: str, middle: str = "") -> str:
    parts = [first, middle, last]
    return " ".join(p for p in parts if p)


async def get_mx_host(domain: str) -> str | None:
    """Get MX host for a domain."""
    try:
        mx_records = await asyncio.to_thread(dns.resolver.resolve, domain, 'MX')
        return str(sorted(mx_records, key=lambda r: r.preference)[0].exchange)
    except:
        return None


async def verify_smtp_port(mx_host: str, email: str, port: int, use_tls: bool = False, timeout: int = 15) -> dict:
    """Verify email via SMTP on specific port."""
    try:
        def _smtp_check():
            server = smtplib.SMTP(timeout=timeout)
            server.connect(mx_host, port)
            if use_tls:
                server.starttls()
            server.ehlo()
            server.mail('')
            code, msg = server.rcpt(email)
            server.quit()
            return code, msg

        loop = asyncio.get_event_loop()
        code, msg = await loop.run_in_executor(None, _smtp_check)

        if code == 250:
            return {"email": email, "valid": True, "reason": f"Mailbox exists (port {port})", "status": "valid"}
        else:
            reason = str(msg).strip()[:80] if msg else f"SMTP code {code}"
            return {"email": email, "valid": False, "reason": reason, "status": "invalid"}
    except Exception as e:
        return {"email": email, "valid": False, "reason": f"Port {port} failed: {str(e)[:50]}", "status": "error"}


async def verify_smtp(email: str, timeout: int = 15) -> dict:
    """Verify email via SMTP - try multiple methods."""
    domain = email.split('@')[1]

    # 1. MX record lookup
    mx_host = await get_mx_host(domain)
    if not mx_host:
        return {"email": email, "valid": False, "reason": "Domain has no mail server", "status": "invalid"}

    # 2. Check if domain is catch-all (accepts all emails)
    catchall = await check_catchall(domain, timeout)
    if catchall["catchall"]:
        return {"email": email, "valid": None, "reason": "Domain is catch-all (accepts all)", "status": "catchall"}

    # 3. Try multiple SMTP methods
    methods = [
        (25, False),   # Port 25, plain
        (587, True),   # Port 587, TLS
    ]

    for port, use_tls in methods:
        result = await verify_smtp_port(mx_host, email, port, use_tls, timeout)
        if result["status"] == "valid":
            return result

        # Check for specific rejection reasons (works for both invalid and error)
        reason_lower = result["reason"].lower()
        # Zoho-specific message
        if "zoho" in reason_lower:
            return {"email": email, "valid": None, "reason": "Zoho blocks verification from dynamic IPs", "status": "unknown"}
        # Dynamic IP or policy rejection - generic message
        if "dynamic" in reason_lower or "policy" in reason_lower:
            return {"email": email, "valid": None, "reason": "Domain blocks verification from this IP", "status": "unknown"}

        # Clear "user unknown" means invalid - but not policy rejections
        if result["status"] == "invalid" and "rejected" not in reason_lower:
            return result

    # All methods failed - return last result
    return result


async def check_catchall(domain: str, timeout: int = 15) -> dict:
    """Check if domain is catch-all (accepts all emails)."""
    import random
    import string
    import smtplib

    # Generate random test emails
    def random_str(length=8):
        return ''.join(random.choices(string.ascii_lowercase + string.digits, k=length))

    mx_host = await get_mx_host(domain)
    if not mx_host:
        return {"domain": domain, "catchall": False, "reason": "No MX records"}

    # Test with random emails directly (no recursion!)
    test_emails = [
        f"test{random_str()}@{domain}",
        f"verify{random_str()}@{domain}",
    ]

    results = []
    for email in test_emails:
        try:
            def _check():
                server = smtplib.SMTP(timeout=timeout)
                server.connect(mx_host, 25)
                server.ehlo()
                server.mail('')
                code, msg = server.rcpt(email)
                server.quit()
                return code == 250
            is_accepted = await asyncio.get_event_loop().run_in_executor(None, _check)
            results.append({"email": email, "accepted": is_accepted})
        except:
            results.append({"email": email, "accepted": False})

    # Catch-all if all accepted
    accepted = sum(1 for r in results if r["accepted"])
    is_catchall = accepted >= len(test_emails)

    return {
        "domain": domain,
        "catchall": is_catchall,
        "tested": test_emails,
        "accepted_count": accepted,
        "total_tested": len(test_emails)
    }


@app.get("/api/generate")
async def generate_emails(
    first: str = Query(...),
    last: str = Query(default=""),
    domain: str = Query(...),
    middle: str = Query(default=""),
):
    """Generate possible email patterns for a person."""
    domain = normalize_domain(domain)
    emails = generate_patterns(first, last, domain, middle)
    return {"emails": emails, "count": len(emails)}


@app.post("/api/generate-multi")
async def generate_emails_multi(payload: MultiNameRequest = Body(...)):
    """Generate possible email patterns for multiple people in one domain."""
    domain = normalize_domain(payload.domain)

    groups = []
    total_count = 0
    for person in payload.names:
        normalized = normalize_person_name(person)
        if not normalized:
            continue

        emails = generate_patterns(
            normalized["first"],
            normalized["last"],
            domain,
            normalized["middle"],
        )
        total_count += len(emails)
        groups.append(
            {
                "name": normalized,
                "label": person_label(
                    normalized["first"], normalized["last"], normalized["middle"]
                ),
                "emails": emails,
                "count": len(emails),
            }
        )

    if not groups:
        raise HTTPException(status_code=400, detail="At least one valid first name is required")

    return {"domain": domain, "groups": groups, "total_count": total_count}


@app.get("/api/verify")
async def verify_email(email: str = Query(...)):
    """Verify a single email address."""
    result = await verify_smtp(email)
    return result


@app.post("/api/verify-all-multi")
async def verify_all_multi(payload: MultiNameRequest = Body(...)):
    """Generate and verify email patterns for multiple people in one domain."""
    domain = normalize_domain(payload.domain)

    groups = []
    verification_jobs = []

    for person in payload.names:
        normalized = normalize_person_name(person)
        if not normalized:
            continue

        emails = generate_patterns(
            normalized["first"],
            normalized["last"],
            domain,
            normalized["middle"],
        )
        group_index = len(groups)
        groups.append(
            {
                "name": normalized,
                "label": person_label(
                    normalized["first"], normalized["last"], normalized["middle"]
                ),
                "results": [],
            }
        )
        for email in emails:
            verification_jobs.append((group_index, verify_smtp(email)))

    if not groups:
        raise HTTPException(status_code=400, detail="At least one valid first name is required")

    if verification_jobs:
        verified_results = await asyncio.gather(*[job for _, job in verification_jobs])
        for (group_index, _), result in zip(verification_jobs, verified_results):
            groups[group_index]["results"].append(result)

    return {"domain": domain, "groups": groups}


@app.get("/api/verify-all")
async def verify_all(
    first: str = Query(...),
    last: str = Query(default=""),
    domain: str = Query(...),
    middle: str = Query(default=""),
    stream: bool = Query(default=False),
):
    """Generate and verify all email patterns. Set stream=true for sequential verification."""
    domain = normalize_domain(domain)
    emails = generate_patterns(first, last, domain, middle)

    if stream:
        # Sequential verification - one at a time
        results = []
        for email in emails:
            result = await verify_smtp(email)
            results.append(result)
        return {"results": results}
    else:
        # Parallel verification (default)
        tasks = [verify_smtp(email) for email in emails]
        results = await asyncio.gather(*tasks)
        return {"results": results}



@app.get("/api/verify-stream")
async def verify_stream(
    first: str = Query(...),
    last: str = Query(default=""),
    domain: str = Query(...),
    middle: str = Query(default=""),
):
    """Stream verification results as they complete (SSE) - sequential verification."""
    domain = normalize_domain(domain)
    emails = generate_patterns(first, last, domain, middle)

    async def event_generator():
        try:
            for email in emails:
                try:
                    result = await verify_smtp(email)
                    yield f"data: {json.dumps(result)}\n\n"
                except Exception as e:
                    logger.error(f"Error verifying {email}: {e}")
                    yield f"data: {json.dumps({'email': email, 'status': 'error', 'valid': False, 'reason': str(e)[:100]})}\n\n"
            # Send completion signal
            yield f"data: {json.dumps({'status': 'done'})}\n\n"
        except Exception as e:
            logger.error(f"Stream error: {e}")
            yield f"data: {json.dumps({'status': 'error', 'error': str(e)[:100]})}\n\n"

    headers = {
        "Cache-Control": "no-cache",
        "X-Accel-Buffering": "no",
        "Connection": "keep-alive",
    }
    return StreamingResponse(event_generator(), media_type="text/event-stream", headers=headers)


@app.post("/api/verify-stream-multi")
async def verify_stream_multi(payload: MultiNameRequest = Body(...)):
    """Stream verification results for multiple people (SSE), sequentially."""
    domain = payload.domain.strip()
    if not domain:
        raise HTTPException(status_code=400, detail="Domain is required")

    normalized_people = []
    for person in payload.names:
        normalized = normalize_person_name(person)
        if normalized:
            normalized_people.append(normalized)

    if not normalized_people:
        raise HTTPException(status_code=400, detail="At least one valid first name is required")

    async def event_generator():
        try:
            for person in normalized_people:
                label = person_label(person["first"], person["last"], person["middle"])
                emails = generate_patterns(
                    person["first"], person["last"], domain, person["middle"]
                )
                yield f"data: {json.dumps({'status': 'person-start', 'label': label, 'name': person})}\n\n"

                for email in emails:
                    try:
                        result = await verify_smtp(email)
                        result["name"] = person
                        result["label"] = label
                        yield f"data: {json.dumps(result)}\n\n"
                    except Exception as e:
                        logger.error(f"Error verifying {email}: {e}")
                        error_result = {
                            "email": email,
                            "status": "error",
                            "valid": False,
                            "reason": str(e)[:100],
                            "name": person,
                            "label": label,
                        }
                        yield f"data: {json.dumps(error_result)}\n\n"

            yield f"data: {json.dumps({'status': 'done'})}\n\n"
        except Exception as e:
            logger.error(f"Multi stream error: {e}")
            yield f"data: {json.dumps({'status': 'error', 'error': str(e)[:100]})}\n\n"

    headers = {
        "Cache-Control": "no-cache",
        "X-Accel-Buffering": "no",
        "Connection": "keep-alive",
    }
    return StreamingResponse(event_generator(), media_type="text/event-stream", headers=headers)

@app.get("/api/check-catchall")
async def api_check_catchall(domain: str = Query(...)):
    """Check if domain is catch-all (accepts all emails)."""
    return await check_catchall(domain)


@app.post("/api/validate-emails")
async def validate_emails(payload: ValidateEmailsRequest = Body(...)):
    """Validate a list of provided emails."""
    if not payload.emails:
        raise HTTPException(status_code=400, detail="At least one email is required")
    
    # Remove duplicates and validate email format
    valid_emails = dedupe_emails([
        email.strip()
        for email in payload.emails
        if email.strip() and EMAIL_RE.fullmatch(email.strip())
    ])
    
    if not valid_emails:
        raise HTTPException(status_code=400, detail="No valid emails provided")
    
    # Verify all emails in parallel
    tasks = [verify_smtp(email) for email in valid_emails]
    results = await asyncio.gather(*tasks)
    
    return {"total_count": len(results), "results": results}


@app.post("/api/validate-emails-upload")
async def validate_emails_upload(
    file: UploadFile | None = File(default=None),
    emails_text: str = Form(default=""),
):
    """Validate emails from pasted text, uploaded CSV/XLS/XLSX, or both."""
    extracted_emails: list[str] = []

    if emails_text.strip():
        extracted_emails.extend(extract_emails_from_text(emails_text))

    if file is not None:
        content = await file.read()
        if not content:
            raise HTTPException(status_code=400, detail="Uploaded file is empty")
        try:
            extracted_emails.extend(extract_emails_from_file(file.filename or "", content))
        except Exception as exc:
            logger.exception("Failed to parse uploaded file")
            raise HTTPException(status_code=400, detail=f"Could not extract emails from file: {str(exc)[:120]}")

    valid_emails = dedupe_emails([
        email for email in extracted_emails if EMAIL_RE.fullmatch(email.strip())
    ])

    if not valid_emails:
        raise HTTPException(status_code=400, detail="No valid emails provided")

    tasks = [verify_smtp(email) for email in valid_emails]
    results = await asyncio.gather(*tasks)

    return {"total_count": len(results), "results": results, "source_count": len(extracted_emails)}


@app.get("/", response_class=HTMLResponse)
async def index():
    html = Path("static/index.html").read_text()
    return HTMLResponse(html)